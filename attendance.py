import cv2
import numpy as np
import pandas as pd
from datetime import datetime
import os
import json
import mediapipe as mp
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ── MediaPipe Face Mesh ──────────────────────────────────────────────────────
mp_face_mesh = mp.solutions.face_mesh
face_mesh = mp_face_mesh.FaceMesh(refine_landmarks=True)

LEFT_EYE  = [33, 160, 158, 133, 153, 144]
RIGHT_EYE = [362, 385, 387, 263, 373, 380]
EAR_THRESHOLD   = 0.25
BLINKS_REQUIRED = 2

def eye_ratio(landmarks, eye_points, w, h):
    pts = [(int(landmarks[i].x * w), int(landmarks[i].y * h)) for i in eye_points]
    v1 = np.linalg.norm(np.array(pts[1]) - np.array(pts[5]))
    v2 = np.linalg.norm(np.array(pts[2]) - np.array(pts[4]))
    h1 = np.linalg.norm(np.array(pts[0]) - np.array(pts[3]))
    return (v1 + v2) / (2.0 * h1)

# ── Load Roll Numbers from JSON ───────────────────────────────────────────────
ROLL_FILE = "roll_numbers.json"
if os.path.exists(ROLL_FILE):
    with open(ROLL_FILE, "r") as f:
        ROLL_NUMBERS = json.load(f)
else:
    ROLL_NUMBERS = {}
    print("Warning: roll_numbers.json not found. Roll numbers will show as N/A.")

# ── Load Face Recognition Model ──────────────────────────────────────────────
model = cv2.face.LBPHFaceRecognizer_create()
model.read("trained_model.xml")
labels = np.load("labels.npy", allow_pickle=True).item()

# ── Face Detector ────────────────────────────────────────────────────────────
face_cascade = cv2.CascadeClassifier(
    cv2.data.haarcascades + 'haarcascade_frontalface_default.xml'
)

CONFIDENCE_THRESHOLD = 60
EXCEL_FILE = "attendance.xlsx"

# ── Excel Setup ──────────────────────────────────────────────────────────────
def setup_excel():
    if not os.path.exists(EXCEL_FILE):
        wb = Workbook()
        ws = wb.active
        ws.title = "Attendance"

        header_font  = Font(name="Arial", bold=True, color="FFFFFF", size=12)
        header_fill  = PatternFill("solid", start_color="2F75B6")
        header_align = Alignment(horizontal="center", vertical="center")
        thin_border  = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"),  bottom=Side(style="thin")
        )

        headers    = ["S.No", "Roll No", "Name", "Date", "Time"]
        col_widths = [8, 12, 25, 18, 15]

        for col, (header, width) in enumerate(zip(headers, col_widths), start=1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font      = header_font
            cell.fill      = header_fill
            cell.alignment = header_align
            cell.border    = thin_border
            ws.column_dimensions[cell.column_letter].width = width

        ws.row_dimensions[1].height = 25
        wb.save(EXCEL_FILE)

def mark_attendance_excel(name, roll_no, date_str, time_str):
    wb = load_workbook(EXCEL_FILE)
    ws = wb["Attendance"]

    next_row  = ws.max_row + 1
    serial_no = next_row - 1

    row_align  = Alignment(horizontal="center")
    name_align = Alignment(horizontal="left")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"),  bottom=Side(style="thin")
    )
    row_fill = PatternFill("solid", start_color="DEEAF1") if serial_no % 2 == 0 else PatternFill("solid", start_color="FFFFFF")

    data   = [serial_no, roll_no, name, date_str, time_str]
    aligns = [row_align, row_align, name_align, row_align, row_align]

    for col, (value, align) in enumerate(zip(data, aligns), start=1):
        cell = ws.cell(row=next_row, column=col, value=value)
        cell.font      = Font(name="Arial", size=11)
        cell.alignment = align
        cell.border    = thin_border
        cell.fill      = row_fill

    wb.save(EXCEL_FILE)

setup_excel()

# ── Per-person blink tracking ─────────────────────────────────────────────────
state = {}

cap = cv2.VideoCapture(0)

while True:
    ret, frame = cap.read()
    if not ret:
        break

    h_frame, w_frame, _ = frame.shape
    gray  = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
    faces = face_cascade.detectMultiScale(gray, 1.3, 5)

    # ── Blink detection ───────────────────────────────────────────────────────
    rgb_frame = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
    results   = face_mesh.process(rgb_frame)

    blink_this_frame = False
    if results.multi_face_landmarks:
        for face_landmarks in results.multi_face_landmarks:
            leftEAR  = eye_ratio(face_landmarks.landmark, LEFT_EYE,  w_frame, h_frame)
            rightEAR = eye_ratio(face_landmarks.landmark, RIGHT_EYE, w_frame, h_frame)
            blink_this_frame = (leftEAR + rightEAR) / 2.0 < EAR_THRESHOLD

    # ── Process each detected face ────────────────────────────────────────────
    for (x, y, w, h) in faces:
        face_roi = gray[y:y+h, x:x+w]
        face_roi = cv2.resize(face_roi, (200, 200))

        label_id, confidence = model.predict(face_roi)

        if confidence >= CONFIDENCE_THRESHOLD:
            cv2.rectangle(frame, (x, y), (x+w, y+h), (0, 0, 255), 2)
            continue

        name    = labels[label_id]
        roll_no = ROLL_NUMBERS.get(name, "N/A")

        if name not in state:
            state[name] = {"blinks": 0, "counter": 0, "marked": False}

        s = state[name]

        if not s["marked"]:
            if blink_this_frame:
                s["counter"] += 1
            else:
                if s["counter"] >= 2:
                    s["blinks"] += 1
                s["counter"] = 0

        if s["blinks"] >= BLINKS_REQUIRED and not s["marked"]:
            s["marked"] = True
            now = datetime.now()
            mark_attendance_excel(name, roll_no, now.strftime("%d-%m-%Y"), now.strftime("%H:%M:%S"))

        if s["marked"]:
            status_text = f"{name} ({roll_no}) - Marked"
            box_color   = (0, 255, 0)
        else:
            blinks_left = BLINKS_REQUIRED - s["blinks"]
            status_text = f"{name} ({roll_no}) - Blink {blinks_left}x"
            box_color   = (255, 165, 0)

        cv2.rectangle(frame, (x, y), (x+w, y+h), box_color, 2)
        cv2.putText(frame, status_text, (x, y - 10),
                    cv2.FONT_HERSHEY_SIMPLEX, 0.8, box_color, 2)

    cv2.imshow("Attendance System", frame)

    if cv2.waitKey(1) == 27:
        break

cap.release()
cv2.destroyAllWindows()